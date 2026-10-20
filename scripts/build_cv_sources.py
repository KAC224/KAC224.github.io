from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).resolve().parents[1]
cv=root/'cv'
source=cv/'Khadijah_Aamir_Chaudhry_Original_CV.docx'
base=Document(source)
paras=base.paragraphs
site='https://kac224.github.io/'
links={
 'CHASE Lab':'https://chase.lums.edu.pk/',
 'Dr. Muhammad Ali Siddiqi':'https://sbasse.lums.edu.pk/lums_employee/9462',
 'SysNet Lab':'https://sysnet.lums.edu.pk/',
 'Dr. Naveed Anwar Bhatti':'https://sbasse.lums.edu.pk/lums_employee/9361',
 'On Chip Behavioral Monitoring':site+'assets/writeups/on-chip-behavioral-monitoring-draft.pdf',
 'Protected Media Containers':site+'assets/writeups/protected-media-containers-early-draft.pdf',
 'EyeSense: Mobile Eyesight Screening':site+'assets/writeups/eyesense.pdf',
 'Post-Quantum Cryptography on ESP32':site+'assets/writeups/post-quantum-esp32.pdf',
 'Fintelligent Solutions':'https://www.fintelligent.cc/',
 'Uplift AI':'https://upliftai.org/',
 'Women in STEM Coordinator':'https://spades.lums.edu.pk/initiatives/women-in-stem',
 'AasPass Peer Tutoring':'https://drive.google.com/drive/folders/1dZr6n9DSSht6m3_ncmZnLsxQ2xPbCkZq',
 'Spotify Mood Classification and Recommendation':'https://medium.com/@khadijah.aamir.ch/spotting-the-vibe-using-machine-learning-techniques-to-predict-song-mood-and-provide-music-6a4544c3de8e',
 'ROM-Based Digital Lock':site+'assets/rom-based-digital-lock-report.pdf',
}

def hyperlink_runs(doc):
 # Keep the original font, emphasis and size, with blue text for every link.
 for p in doc.paragraphs:
  for label,url in links.items():
   for r in list(p._p.findall(qn('w:r'))):
    text=''.join(t.text or '' for t in r.findall(qn('w:t')))
    if label not in text: continue
    # All linked labels occur within a single text run in the original template.
    before,after=text.split(label,1)
    new=[]
    def fragment(text):
     rr=OxmlElement('w:r')
     props=r.find(qn('w:rPr'))
     if props is not None: rr.append(deepcopy(props))
     tt=OxmlElement('w:t');tt.set(qn('xml:space'),'preserve');tt.text=text;rr.append(tt)
     return rr
    if before:new.append(fragment(before))
    h=OxmlElement('w:hyperlink')
    h.set(qn('r:id'),doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
    h.append(fragment(label));new.append(h)
    if after:new.append(fragment(after))
    idx=p._p.index(r);p._p.remove(r)
    for el in new:p._p.insert(idx,el);idx+=1
  for r in p._p.xpath('.//w:hyperlink/w:r'):
   props=r.find(qn('w:rPr'))
   if props is None:
    props=OxmlElement('w:rPr');r.insert(0,props)
   color=props.find(qn('w:color'))
   if color is None:
    color=OxmlElement('w:color');props.append(color)
   color.attrib.clear()
   color.set(qn('w:val'),'175A86')

pro=Document(source)
hyperlink_runs(pro)
pro.core_properties.title='Khadijah Aamir Chaudhry — Professional CV'
pro.save(cv/'Khadijah_Aamir_Chaudhry_Professional_CV.docx')

academic=Document(source)
# Retain every original font, size, border and horizontal alignment.
# Increase only vertical spacing for the academic version.
for name in ['Normal','CV Entry','CV Bullet','CV Credential','CV Advisor','CV Skills']:
 academic.styles[name].paragraph_format.line_spacing=1.12
academic.styles['CV Section'].paragraph_format.space_before=Pt(9)
academic.styles['CV Section'].paragraph_format.space_after=Pt(4)
academic.styles['CV Entry'].paragraph_format.space_before=Pt(6)
academic.styles['CV Bullet'].paragraph_format.space_after=Pt(4)
academic.styles['CV Advisor'].paragraph_format.space_after=Pt(2)

# Each tuple is (action-led bullet text, one selectively bold phrase).
expanded={
13:[
 ('Benchmarked post-quantum algorithms on ESP32 to compare latency, memory, network cost, and estimated energy on constrained devices.','post-quantum algorithms on ESP32'),
 ('Built a simulated RISC-V trace testbed to capture transaction-level activity for research into runtime hardware Trojan detection.','simulated RISC-V trace testbed')],
16:[
 ('Co-developed EyeSense’s smartphone ToF/ML pipeline for refractive-error screening to explore more accessible eyesight measurement.','EyeSense’s smartphone ToF/ML pipeline'),
 ('Helped collect eyesight data in lab and eye-camp settings, supporting evaluation of phone-based screening with real measurements.','eyesight data in lab and eye-camp settings')],
19:[
 ('Built a Rocket Chip/Verilator testbed with four RV32 configurations and TileLink MMIO trace capture for runtime hardware monitoring research.','TileLink MMIO trace capture'),
 ('Replayed benign and trigger-active firmware to collect contrasting traces for later detector evaluation.','benign and trigger-active firmware')],
21:[
 ('Developed a system design for protected camera/audio capture, encrypted storage, display, and sharing under an untrusted mobile OS.','protected camera/audio capture'),
 ('Reviewed trusted-execution and DRM research to define how privacy controls can persist as media is stored and shared.','trusted-execution and DRM research')],
23:[
 ('Co-developed a smartphone ToF phase-map and machine-learning pipeline for refractive-error screening, exploring a phone-based approach to eyesight checks.','smartphone ToF phase-map and machine-learning pipeline'),
 ('Supported a team study of 120 participants that reported 1.15 D mean absolute error, providing an initial measure of screening performance.','1.15 D mean absolute error')],
25:[
 ('Benchmarked Kyber-512, Light-SABER, and Dilithium-II on ESP32 to assess post-quantum cryptography on resource-constrained hardware.','Kyber-512, Light-SABER, and Dilithium-II'),
 ('Compared latency, memory, network cost, and estimated energy use to document practical trade-offs for small connected devices.','latency, memory, network cost, and estimated energy use')],
28:[
 ('Designed and conducted programming labs and prepared and graded quizzes to support about 450 first-year students learning introductory programming.','450 first-year students'),
 ('Held weekly office hours and invigilated exams, providing regular opportunities for students to seek help and receive assessment support.','weekly office hours')],
31:[
 ('Coordinated team delivery of brand identity and the company website for a Microsoft Dynamics 365 consultancy, helping clients understand its services.','brand identity and the company website'),
 ('Coordinated office collateral and LinkedIn launch content to give the consultancy consistent materials for client outreach.','office collateral and LinkedIn launch content')],
33:[
 ('Led 10 data collectors and organized field outreach for a Pakistani-language AI dataset, helping gather locally relevant voice data.','10 data collectors'),
 ('Ran daily quality reviews and feedback sessions to identify collection issues and improve consistency in submitted recordings.','daily quality reviews and feedback sessions')],
36:[
 ('Guided 15 incoming students through orientation and faculty introductions to help them settle into university life.','15 incoming students'),
 ('Conduct monthly first-year support sessions to give students a familiar point of contact as they navigate university.','monthly first-year support sessions')],
38:[
 ('Helped launch the first women-only STEM event at PSIFI for more than 100 participants, creating a dedicated space for women to engage with STEM.','first women-only STEM event at PSIFI'),
 ('Led the Women in STEM department in organizing bonding events and digital and financial literacy seminars to support participation and practical learning.','digital and financial literacy seminars')],
40:[
 ('Co-founded Umeed-e-Nau, an NGO providing inclusive educational activities for underprivileged and disabled children.','inclusive educational activities'),
 ('Directed logistics for tailored learning activities and visits, including accommodations that helped disabled children participate.','tailored learning activities and visits')],
}

def bullet(text,bold):
 p=deepcopy(paras[13]._p)
 for child in list(p):
  if child.tag!=qn('w:pPr'): p.remove(child)
 from docx.text.paragraph import Paragraph
 obj=Paragraph(p,academic._body)
 before,after=text.split(bold,1)
 obj.add_run('• '+before)
 obj.add_run(bold).bold=True
 obj.add_run(after)
 obj.paragraph_format.keep_together=True
 return obj

def heading(title,date=''):
 p=deepcopy(paras[18]._p)
 for c in list(p):
  if c.tag!=qn('w:pPr'):p.remove(c)
 from docx.text.paragraph import Paragraph
 obj=Paragraph(p,academic._body)
 obj.add_run(title).bold=True
 if date:
  obj.add_run('\t');obj.add_run(date).italic=True
 return obj

extras=[
 ('AasPass Peer Tutoring','',[
 ('Led authentication and administration features for an Android student–tutor app, including LUMS email verification and profile setup.','LUMS email verification and profile setup'),
 ('Added moderation controls so administrators could manage accounts and support the peer tutoring community.','moderation controls')]),
 ('Spotify Mood Classification and Recommendation','Jan–May 2025',[
 ('Applied PCA and K-means to audio features to explore how songs group by mood and musical similarity.','PCA and K-means'),
 ('Built mood classification and similarity-based recommendations to help listeners discover songs with a similar feel.','mood classification and similarity-based recommendations')]),
 ('ROM-Based Digital Lock','',[
 ('Built and tested an EPROM-based keypad lock using state-machine logic to control access based on an entered code.','EPROM-based keypad lock'),
 ('Added an unlock indicator and an incorrect-input alarm to make accepted and rejected codes clear to users.','unlock indicator and an incorrect-input alarm')]),
]
for i,p in enumerate(list(academic.paragraphs)):
 if i in expanded:
  for j,(text,bold) in enumerate(expanded[i]):
   b=bullet(text,bold)
   b.paragraph_format.keep_with_next=(j==0)
   p._p.addprevious(b._p)
  p._p.getparent().remove(p._p)
 if i==26:
  for title,date,points in extras:
   p._p.addprevious(heading(title,date)._p)
   for j,(text,bold) in enumerate(points):
    b=bullet(text,bold);b.paragraph_format.keep_with_next=(j==0)
    p._p.addprevious(b._p)

hyperlink_runs(academic)
academic.core_properties.title='Khadijah Aamir Chaudhry — Academic CV'
academic.save(cv/'Khadijah_Aamir_Chaudhry_Academic_CV.docx')
(root/'scripts'/'cv-links.json').write_text(json.dumps(links,ensure_ascii=False,indent=2)+'\n')
print('Created both Word documents from the pre-redesign original.')
